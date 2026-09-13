import re
from collections import defaultdict
from pathlib import Path
import fitz
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = Path('/mnt/data')

PAIR_CONFIGS = [
    {
        'pair': 'IPC ↔ BNS',
        'comparison_pdf': BASE/'COMPARISON_SUMMARY_IPC_to_BNS.pdf',
        'thresholds': (100, 301, 349),
        'order': 'new-subj-old-sum',
        'act_old': 'Indian Penal Code, 1860 (IPC)',
        'act_new': 'Bharatiya Nyaya Sanhita, 2023 (BNS)',
        'old_pdf': BASE/'IPC.pdf',
        'new_pdf': BASE/'BNS.pdf',
        'old_start_page': 14,
        'new_start_page': 1,
        'comparison_url': 'https://indialawacts.in/comparisonIPCtoBNS.php',
    },
    {
        'pair': 'CrPC ↔ BNSS',
        'comparison_pdf': BASE/'COMPARISON_SUMMARY_CrPC_to_BNSS.pdf',
        'thresholds': (80, 239, 285),
        'order': 'new-subj-old-sum',
        'act_old': 'Code of Criminal Procedure, 1973 (CrPC)',
        'act_new': 'Bharatiya Nagarik Suraksha Sanhita, 2023 (BNSS)',
        'old_pdf': BASE/'CrPC.pdf',
        'new_pdf': BASE/'BNSS.pdf',
        'old_start_page': 21,
        'new_start_page': 16,
        'comparison_url': 'https://indialawacts.in/comparisonCrPCtoBNSS.php',
    },
    {
        'pair': 'IEA ↔ BSA',
        'comparison_pdf': BASE/'COMPARISON_SUMMARY_IEA_to_BSA.pdf',
        'thresholds': (90, 146, 317),
        'order': 'new-old-subj-sum',
        'act_old': 'Indian Evidence Act, 1872 (IEA)',
        'act_new': 'Bharatiya Sakshya Adhiniyam, 2023 (BSA)',
        'old_pdf': BASE/'IEA.pdf',
        'new_pdf': BASE/'BSA.pdf',
        'old_start_page': 9,
        'new_start_page': 10,
        'comparison_url': 'https://indialawacts.in/comparisonIEAtoBSA.php',
    },
]

# Optional authoritative URLs for source notes
OLD_NEW_OFFICIAL_URLS = {
    'Bharatiya Nyaya Sanhita, 2023 (BNS)': 'https://www.indiacode.nic.in/bitstream/123456789/20062/1/a202345.pdf',
    'Bharatiya Nagarik Suraksha Sanhita, 2023 (BNSS)': 'https://www.indiacode.nic.in/bitstream/123456789/21544/1/the_bharatiya_nagarik_suraksha_sanhita%2C_2023.pdf',
    'Bharatiya Sakshya Adhiniyam, 2023 (BSA)': 'https://upload.indiacode.nic.in/view-casepdf?id=AC_CEN_5_23_00049_2023-47_1719292804654&type=act',
}


def clean_space(s: str) -> str:
    return re.sub(r'\s+', ' ', (s or '')).strip(' |\t\n')


def clean_ref(ref: str) -> str:
    ref = clean_space(ref)
    ref = ref.replace(' ]', ']').replace('[ ', '[')
    ref = ref.replace(' )', ')').replace('( ', '(')
    ref = re.sub(r'\s+', ' ', ref)
    ref = ref.replace(' ,', ',')
    ref = re.sub(r'(\d)\s+\(', r'\1(', ref)
    ref = re.sub(r'\)\s+\(', ')(', ref)
    ref = re.sub(r'\[\s*', '[', ref)
    ref = re.sub(r'\s*\]', ']', ref)
    ref = ref.rstrip('-').strip()
    # clean some OCR-ish closers
    if ref.endswith(')]') and ref.count('(') >= ref.count(')'):
        ref = ref[:-1]
    return ref


def normalize_bracket_text(subject: str) -> str:
    s = clean_space(subject)
    s = s.strip('"“”')
    s = s.strip('. ')
    return s


def base_section_from_ref(ref: str) -> str:
    ref = clean_ref(ref)
    if not ref or ref in {'New', '-', '—'}:
        return ''
    m = re.search(r'(\d+[A-Z]?)', ref)
    return m.group(1) if m else ''


def extract_table_lines(pdf_path: Path, thresholds):
    doc = fitz.open(pdf_path)
    lines = []
    for pno in range(doc.page_count):
        words = doc[pno].get_text('words')
        groups = defaultdict(list)
        for w in words:
            x0, y0, x1, y1, text, *_ = w
            groups[round(y0, 1)].append((x0, text))
        for y in sorted(groups):
            cols = ['', '', '', '']
            for x0, text in sorted(groups[y]):
                if x0 >= thresholds[2]:
                    idx = 3
                elif x0 >= thresholds[1]:
                    idx = 2
                elif x0 >= thresholds[0]:
                    idx = 1
                else:
                    idx = 0
                cols[idx] += ('' if not cols[idx] else ' ') + text
            lines.append((pno + 1, y, *cols))
    return lines


def is_data_start(c1: str) -> bool:
    c1 = c1.strip()
    return bool(re.match(r'^(?:\d+[A-Z]?(?:\s*\([^)]+\))*(?:\s*,\s*para\s*\d+)?|First proviso|Second proviso|Third proviso|Fourth proviso|Proviso|Explanation|Exception|Illustration)', c1, re.I))


def parse_comparison_pdf(pdf_path: Path, thresholds, order: str):
    lines = extract_table_lines(pdf_path, thresholds)
    rows = []
    curr = None
    for p, y, c1, c2, c3, c4 in lines:
        joined = ' '.join([c1, c2, c3, c4]).strip()
        if not joined:
            continue
        # Skip headers and boilerplate
        lower = joined.lower()
        if any(x in lower for x in [
            'correspondence table', 'comparison summary',
            'bhartatiya', 'bharatiya', 'indian penal code',
            'code of criminal procedure', 'indian evidence act',
            'summary of comparison', 'summary of comparision'
        ]):
            continue
        if re.fullmatch(r'\d+', joined):
            continue
        if is_data_start(c1):
            if curr:
                rows.append(curr)
            curr = [p, c1, c2, c3, c4]
        elif curr:
            for i, v in enumerate([c1, c2, c3, c4], start=1):
                if v:
                    curr[i] = (curr[i] + ' ' + v).strip()
    if curr:
        rows.append(curr)

    out = []
    for p, c1, c2, c3, c4 in rows:
        c1, c2, c3, c4 = [clean_space(x) for x in (c1, c2, c3, c4)]
        if order == 'new-old-subj-sum':
            new_id, old_ref, subject, summary = c1, c2, c3, c4
        else:
            new_id, subject, old_ref, summary = c1, c2, c3, c4
        # minor cleanup
        new_id = clean_ref(new_id)
        old_ref = clean_ref(old_ref)
        if new_id in {'BNSS', 'BNS', 'BSA'} or subject.lower() == 'subject':
            continue
        out.append({
            'page': p,
            'new_id': new_id,
            'old_ref': old_ref,
            'subject': clean_space(subject),
            'summary': clean_space(summary),
        })
    return out


def build_sections(pdf_path: Path, start_page: int):
    doc = fitz.open(pdf_path)
    txt = '\n'.join(doc[i].get_text('text') for i in range(start_page - 1, doc.page_count))
    txt = txt.replace('\xa0', ' ')
    # section number + period + whitespace + title/body starting with capital/opening punctuation
    pat = re.compile(r'(?m)^\s*(\d+[A-Z]?)\.\s+(?=[A-Z"(\[])')
    matches = list(pat.finditer(txt))
    sections = {}
    for i, m in enumerate(matches):
        sec = m.group(1)
        start = m.start()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(txt)
        body = clean_space(txt[start:end])
        sections.setdefault(sec, body)
    return sections


def strip_leading_section_number(text: str) -> str:
    if not text:
        return ''
    return re.sub(r'^\s*\d+[A-Z]?\.\s*', '', text).strip()


def extract_subunit(section_text: str, exact_ref: str) -> str:
    """Best-effort extraction for sub-sections / clauses. Falls back to section text without number."""
    if not section_text:
        return ''
    text = strip_leading_section_number(section_text)
    ref = clean_ref(exact_ref)
    if not ref or ref in {'New', '-', '—'}:
        return ''
    # para references are hard to isolate reliably in IEA from PDF text; keep base section text.
    if 'para' in ref.lower():
        return text
    # Subsection/letter extraction based on last parenthesized token
    toks = re.findall(r'\(([^)]+)\)', ref)
    if toks:
        last_tok = toks[-1].strip()
        marker = f'({last_tok})'
        idx = text.find(marker)
        if idx != -1:
            content = text[idx:]
            candidates = []
            if last_tok.isdigit():
                n = int(last_tok)
                for k in range(n + 1, min(n + 10, n + 40)):
                    j = content.find(f'({k})', len(marker))
                    if j != -1:
                        candidates.append(j)
            elif len(last_tok) == 1 and last_tok.isalpha():
                start_ord = ord(last_tok.lower())
                for k in range(start_ord + 1, min(start_ord + 8, ord('z') + 1)):
                    j = content.find(f'({chr(k)})', len(marker))
                    if j != -1:
                        candidates.append(j)
            # generic breakpoints
            for lab in ['Explanation', 'Illustration', 'Provided that', 'Provided further that', 'Exception']:
                j = content.find(lab, len(marker))
                if j != -1:
                    candidates.append(j)
            end = min(candidates) if candidates else len(content)
            snippet = clean_space(content[:end])
            return snippet
    return text


def build_rows():
    act_text_cache = {}
    all_rows = []
    for cfg in PAIR_CONFIGS:
        old_sections = build_sections(cfg['old_pdf'], cfg['old_start_page'])
        new_sections = build_sections(cfg['new_pdf'], cfg['new_start_page'])
        comp_rows = parse_comparison_pdf(cfg['comparison_pdf'], cfg['thresholds'], cfg['order'])

        for r in comp_rows:
            old_ref = r['old_ref']
            new_ref = r['new_id']
            old_base = base_section_from_ref(old_ref)
            new_base = base_section_from_ref(new_ref)
            old_sec_text = old_sections.get(old_base, '') if old_base else ''
            new_sec_text = new_sections.get(new_base, '') if new_base else ''
            bracket = normalize_bracket_text(r['subject'])
            old_bracket = '' if old_ref in {'', 'New', '-', '—'} else bracket
            new_bracket = '' if new_ref in {'', 'New', '-', '—'} else bracket
            all_rows.append({
                'Pair': cfg['pair'],
                'Act_old': cfg['act_old'],
                'section number_old': old_ref,
                'Bracket Text_old': old_bracket,
                'clause_old': extract_subunit(old_sec_text, old_ref),
                'Description_old': '',
                'Act_new': cfg['act_new'],
                'section number_new': new_ref,
                'Bracket Text_new': new_bracket,
                'clause_new': extract_subunit(new_sec_text, new_ref),
                'Description_new': '',
                'Comparison Summary': r['summary'],
                'Comparison_Source_URL': cfg['comparison_url'],
                'Old_Act_PDF_URL': '',
                'New_Act_PDF_URL': OLD_NEW_OFFICIAL_URLS.get(cfg['act_new'], ''),
            })
    return all_rows


def style_ws(ws, table_name='Table1'):
    # Freeze and filter
    ws.freeze_panes = 'A2'
    if ws.max_row >= 2 and ws.max_column >= 2:
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        tab = Table(displayName=table_name, ref=ref)
        style = TableStyleInfo(name='TableStyleMedium2', showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        tab.tableStyleInfo = style
        ws.add_table(tab)

    # Header style
    header_fill = PatternFill('solid', fgColor='1F4E78')
    header_font = Font(color='FFFFFF', bold=True)
    thin = Side(style='thin', color='D9E2F3')
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(bottom=thin)

    # Column widths and wrap
    widths = {
        'A': 16, 'B': 32, 'C': 18, 'D': 34, 'E': 70, 'F': 24,
        'G': 36, 'H': 18, 'I': 34, 'J': 70, 'K': 24, 'L': 60,
        'M': 42, 'N': 42, 'O': 42,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical='top', wrap_text=True)
    ws.sheet_view.showGridLines = False


def build_workbook(output_path: Path):
    rows = build_rows()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Mapped_Sections'
    headers = [
        'Pair',
        'Act_old', 'section number_old', 'Bracket Text_old', 'clause_old', 'Description_old',
        'Act_new', 'section number_new', 'Bracket Text_new', 'clause_new', 'Description_new',
        'Comparison Summary', 'Comparison_Source_URL', 'Old_Act_PDF_URL', 'New_Act_PDF_URL'
    ]
    ws.append(headers)
    for row in rows:
        ws.append([row[h] for h in headers])
    style_ws(ws, table_name='MappedSections')

    # Pair-specific sheets
    pair_rows = defaultdict(list)
    for r in rows:
        pair_rows[r['Pair']].append(r)
    for idx, (pair, rws) in enumerate(pair_rows.items(), start=1):
        w = wb.create_sheet(title=pair.replace('↔', '_').replace(' ', '')[:31])
        w.append(headers)
        for r in rws:
            w.append([r[h] for h in headers])
        style_ws(w, table_name=f'PairTable{idx}')

    # Notes sheet
    notes = wb.create_sheet('Notes')
    notes.append(['Item', 'Value'])
    note_rows = [
        ['Scope', 'Section-level correspondence extracted for the three mapped replacement-law pairs only: IPC↔BNS, CrPC↔BNSS, IEA↔BSA.'],
        ['Method', 'Parsed the site comparison-summary PDFs and linked each row to best-effort bare-act text from the corresponding Act PDFs.'],
        ['Bracket Text columns', 'Filled from the comparison-table subject text.'],
        ['clause_old / clause_new', 'Best-effort snippet from the corresponding old/new Act PDF. For subsection-style references, the script tries to isolate the relevant sub-clause; otherwise it falls back to the base section text.'],
        ['Description_old / Description_new', 'Left blank because the comparison PDFs do not expose separate old/new description fields in a machine-clean way, and the site dropdown HTML could not be programmatically exported from this runtime.'],
        ['Comparison Summary', 'Directly from the comparison summary table for that mapping row.'],
        ['Source caution', 'indialawacts.in itself says users should verify with official sources such as India Code for authoritative legal use.'],
    ]
    for r in note_rows:
        notes.append(r)
    notes.column_dimensions['A'].width = 24
    notes.column_dimensions['B'].width = 120
    for row in notes.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(vertical='top', wrap_text=True)
    for c in notes[1]:
        c.fill = PatternFill('solid', fgColor='1F4E78')
        c.font = Font(color='FFFFFF', bold=True)
    notes.sheet_view.showGridLines = False

    wb.save(output_path)
    return len(rows)


if __name__ == '__main__':
    out = BASE/'indian_section_mappings_old_new_pyrun.xlsx'
    n = build_workbook(out)
    print(f'Wrote {out} with {n} mapped rows')
